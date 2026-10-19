from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\akhse\Desktop\REPO\TestLineas")
OUTPUT = ROOT / "docs" / "Metodologia_Scrum_Test_de_Lineas_Base.docx"
SOURCE = Path(r"C:\Users\akhse\Downloads\Metodologia SCRUM - Modelo 1 (1).docx")
document = Document(OUTPUT)
model = Document(SOURCE)
body = document.element.body


def set_text(element, value):
    texts = element.xpath(".//w:t")
    if not texts:
        raise ValueError("No text node in model element")
    texts[0].text = value
    for text in texts[1:]:
        text.text = ""


def replace_paragraph(needle, value):
    matches = [paragraph for paragraph in document.paragraphs if needle in paragraph.text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph for {needle!r}: {len(matches)}")
    set_text(matches[0]._element, value)


replace_paragraph(
    "Este documento presenta la aplicación",
    "Este documento presenta la aplicación de la metodología Scrum al desarrollo de Test de Líneas, una plataforma web para administrar una evaluación visual de nueve estímulos. El paciente completa la prueba mediante un enlace y el profesional responsable consulta los resultados. El área Tesis registra los momentos del proceso evaluativo y permite analizar sus tiempos.",
)
replace_paragraph(
    "Permitir al paciente aceptar el consentimiento",
    "Permitir al paciente aceptar el consentimiento, registrar sus datos y ordenar los nueve estímulos mediante selección o arrastre.",
)
replace_paragraph(
    "Calcular y conservar resultados, tiempos e historial",
    "Calcular y conservar resultados, tiempos de respuesta e historial para el seguimiento profesional.",
)
replace_paragraph(
    "Generar informes, exportar datos e importar pruebas",
    "Generar informes, exportar datos, registrar los tiempos del proceso evaluativo para la tesis e importar pruebas de la plantilla PDF con revisión de los valores detectados.",
)
replace_paragraph(
    "El Product Backlog reúne las capacidades",
    "El Product Backlog reúne las capacidades identificadas en el repositorio de Test de Líneas. Los requerimientos incluyen el flujo actual de consentimiento, ordenamiento de estímulos y registro de tiempos de la sección Tesis. Los tiempos estimados y reales se completarán durante la planificación y ejecución.",
)
replace_paragraph(
    "En la tabla 3 se presentan los requerimientos funcionales",
    "En la tabla 3 se presentan los requerimientos funcionales identificados en el proyecto. Los tiempos estimados y reales quedan vacíos para registrarlos cuando se disponga de esa información.",
)

requirements = [
    ("RF-01", "El sistema debe permitir al profesional iniciar sesión con correo y contraseña para acceder a las funciones restringidas."),
    ("RF-02", "El sistema debe permitir registrar, buscar y actualizar fichas de pacientes y evitar duplicados por documento."),
    ("RF-03", "El sistema debe generar enlaces de evaluación asociados a un paciente, válidos por 24 horas y de un solo uso."),
    ("RF-04", "El sistema debe solicitar el consentimiento, registrarlo antes de continuar y recoger los datos del participante antes de iniciar la prueba."),
    ("RF-05", "El sistema debe mostrar nueve estímulos y permitir ordenarlos sin duplicados mediante selección o arrastre, incluidos los cambios de posición, antes del envío."),
    ("RF-06", "El sistema debe calcular rankings, interpretaciones y tiempos de respuesta a partir del orden enviado."),
    ("RF-07", "El sistema debe guardar cada evaluación asociada al paciente y al profesional y marcar el enlace como completado."),
    ("RF-08", "El sistema debe permitir al profesional consultar y buscar las evaluaciones registradas."),
    ("RF-09", "El sistema debe mostrar el historial y la comparación de evaluaciones de cada paciente."),
    ("RF-10", "El sistema debe mostrar un informe imprimible que pueda guardarse como PDF."),
    ("RF-11", "El sistema debe exportar resultados e historial del paciente a Excel."),
    ("RF-12", "El sistema debe analizar una prueba en PDF, permitir corregir los nueve valores detectados y guardar la evaluación importada."),
    ("RF-13", "El sistema debe registrar en el servidor la hora de aceptación del consentimiento como inicio del proceso evaluativo."),
    ("RF-14", "El sistema debe registrar la hora de confirmación de respuestas y la hora en que queda disponible el resultado de la evaluación."),
    ("RF-15", "El sistema debe mostrar en la sección Tesis las evaluaciones en línea, los tres momentos registrados y los tiempos total del proceso y de procesamiento del resultado, sin inventar datos históricos ausentes."),
    ("RF-16", "El sistema debe exportar a Excel los registros de Tesis con horas de Lima, duraciones en segundos y observaciones cuando falten datos o existan horas inconsistentes."),
]
priorities = {code: ("M" if int(code[-2:]) in {1, 2, 3, 4, 5, 6, 7, 13, 14, 15} else "S") for code, _ in requirements}

backlog = document.tables[2]._element
rows = backlog.xpath("./w:tr")
for _ in range(len(requirements) + 1 - len(rows)):
    backlog.append(deepcopy(rows[-1]))
rows = backlog.xpath("./w:tr")
for row, (code, description) in zip(rows[1:], requirements):
    for cell, value in zip(row.xpath("./w:tc"), (code, description, "", "")):
        set_text(cell, value)

heading = next(p._element for p in document.paragraphs if p.text == "1.4 Product Backlog priorizado")
children = list(body)
start = children.index(heading) + 1
end = next(i for i in range(start, len(children)) if children[i].tag == qn("w:p") and children[i].xpath("./w:pPr/w:sectPr"))
anchor = children[end]
for element in children[start:end]:
    body.remove(element)


def add_paragraph(model_index, value):
    element = deepcopy(model.paragraphs[model_index]._element)
    set_text(element, value)
    body.insert(body.index(anchor), element)


def add_table(element):
    body.insert(body.index(anchor), element)


add_paragraph(399, "El Product Backlog priorizado ordena los requerimientos según su valor para la evaluación, la seguridad de los datos y el análisis de tiempos de la tesis. Esta prioridad es preliminar y debe validarla el Product Owner antes de planificar un Sprint.")
add_paragraph(400, "Prioridad")
add_paragraph(402, "Los requerimientos necesarios para completar la evaluación, conservar sus resultados y medir el proceso evaluativo reciben prioridad M. Las funciones de seguimiento, informes, importación y exportación reciben prioridad S.")
add_paragraph(403, "Técnica MoSCoW:")
add_paragraph(405, "La técnica MoSCoW clasifica los requerimientos según su necesidad para la primera entrega. La tabla 4 conserva las categorías del modelo de referencia.")
add_paragraph(408, "Tabla 4. Prioridad según técnica MoSCoW.")
add_table(deepcopy(model.tables[5]._element))
add_paragraph(410, "Fuente: Elaboración propia.")
add_paragraph(412, "En la tabla 5 se presenta el Product Backlog del proyecto ordenado por prioridad. Cada historia identifica un requerimiento funcional; sus criterios de aceptación y estimación se detallarán durante el refinamiento.")
add_paragraph(414, "Tabla 5. Product backlog priorizado.")

priority_table = deepcopy(model.tables[6]._element)
priority_rows = priority_table.xpath("./w:tr")
for _ in range(len(requirements) + 1 - len(priority_rows)):
    priority_table.append(deepcopy(priority_rows[-1]))
priority_rows = priority_table.xpath("./w:tr")
ordered = sorted(requirements, key=lambda item: (priorities[item[0]] != "M", int(item[0][-2:])))
for index, (row, (code, description)) in enumerate(zip(priority_rows[1:], ordered), 1):
    cells = row.xpath("./w:tc")
    for cell, value in zip(cells, (str(index), f"H.U.{index}", f"{code}: {description}", priorities[code])):
        set_text(cell, value)
add_table(priority_table)
add_paragraph(418, "Fuente: Elaboración propia.")

document.save(OUTPUT)
print(OUTPUT)
