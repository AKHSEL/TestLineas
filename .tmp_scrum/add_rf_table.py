from copy import deepcopy
from pathlib import Path
from docx import Document

reference = Document(r"C:\Users\akhse\Downloads\Metodologia SCRUM - Modelo 1 (1).docx")
output_path = Path(r"C:\Users\akhse\Desktop\REPO\TestLineas\docs\Metodologia_Scrum_Test_de_Lineas_Base.docx")
output = Document(output_path)
body = output.element.body
paragraphs = list(output.paragraphs)
first_bullet = next(p._element for p in paragraphs if p.text == "Acceso profesional y protección de datos por propietario.")
next_heading = next(p._element for p in paragraphs if p.text == "1.4 Product Backlog priorizado")
children = list(body)
for element in children[children.index(first_bullet):children.index(next_heading)]:
    body.remove(element)

def set_text(element, text):
    nodes = element.xpath(".//w:t")
    if not nodes:
        raise ValueError("No text node")
    nodes[0].text = text
    for node in nodes[1:]:
        node.text = ""

def insert_paragraph(index, text=None):
    element = deepcopy(reference.paragraphs[index]._element)
    if text is not None:
        set_text(element, text)
    body.insert(body.index(next_heading), element)

requirements = [
    ("RF-01", "El sistema debe permitir al profesional iniciar sesión con correo y contraseña para acceder a las funciones restringidas."),
    ("RF-02", "El sistema debe permitir registrar, buscar y actualizar fichas de pacientes y evitar duplicados por documento."),
    ("RF-03", "El sistema debe generar enlaces de evaluación asociados a un paciente, válidos por 24 horas y de un solo uso."),
    ("RF-04", "El sistema debe solicitar consentimiento y registrar datos personales antes de iniciar la prueba."),
    ("RF-05", "El sistema debe mostrar nueve estímulos y permitir ordenarlos sin duplicados antes del envío."),
    ("RF-06", "El sistema debe calcular rankings, interpretaciones y tiempos de respuesta a partir del orden enviado."),
    ("RF-07", "El sistema debe guardar cada evaluación asociada al paciente y al profesional y marcar el enlace como completado."),
    ("RF-08", "El sistema debe permitir al profesional consultar y buscar las evaluaciones registradas."),
    ("RF-09", "El sistema debe mostrar el historial y la comparación de evaluaciones de cada paciente."),
    ("RF-10", "El sistema debe mostrar un informe imprimible que pueda guardarse como PDF."),
    ("RF-11", "El sistema debe exportar resultados e historial del paciente a Excel."),
    ("RF-12", "El sistema debe analizar una prueba en PDF, permitir corregir los nueve valores detectados y guardar la evaluación importada."),
]

insert_paragraph(391, "En la tabla 3 se presentan los requerimientos funcionales identificados en el proyecto. Los tiempos estimados y reales quedan vacíos para registrarlos cuando se disponga de esa información.")
insert_paragraph(392)
table = deepcopy(reference.tables[2]._element)
for _ in range(len(requirements) - (len(table.xpath("./w:tr")) - 1)):
    table.append(deepcopy(table.xpath("./w:tr")[-1]))
rows = table.xpath("./w:tr")
for row, (code, description) in zip(rows[1:], requirements):
    cells = row.xpath("./w:tc")
    for cell, value in zip(cells, (code, description, "", "")):
        set_text(cell, value)
body.insert(body.index(next_heading), table)
insert_paragraph(395)
output.save(output_path)
print(output_path)

