from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

root = Path(r"C:\Users\akhse\Desktop\REPO\TestLineas")
reference = Document(r"C:\Users\akhse\Downloads\Metodologia SCRUM - Modelo 1 (1).docx")
output_path = root / "docs" / "Metodologia_Scrum_Test_de_Lineas_Base.docx"
output = Document(output_path)
source = list(reference.paragraphs)
paragraphs = list(output.paragraphs)
start = next(p._element for p in paragraphs if p.text == "1 Alcance del producto")
body = output.element.body
children = list(body)
start_at = children.index(start)
end_at = next(i for i in range(start_at + 1, len(children))
              if children[i].tag == qn("w:p") and children[i].xpath("./w:pPr/w:sectPr"))
end = children[end_at]
for element in children[start_at:end_at]:
    body.remove(element)

def add(source_index, value):
    element = deepcopy(source[source_index]._element)
    texts = element.xpath(".//w:t")
    if texts:
        texts[0].text = value
        for item in texts[1:]:
            item.text = ""
    else:
        run = element.makeelement(qn("w:r"))
        text = run.makeelement(qn("w:t"))
        text.text = value
        run.append(text)
        element.append(run)
    body.insert(body.index(end), element)

add(347, "1.1 Introducción")
add(349, "Este documento presenta la aplicación de la metodología Scrum al desarrollo de Test de Líneas, una plataforma web para administrar una evaluación visual de nueve estímulos. El paciente completa la prueba mediante un enlace y el profesional responsable consulta los resultados para su interpretación.")
add(351, "Alcance")
add(354, "Administrar pacientes, invitaciones y evaluaciones desde un área profesional.")
add(354, "Permitir al paciente aceptar el consentimiento, registrar sus datos y ordenar los nueve estímulos.")
add(354, "Calcular y conservar resultados, tiempos e historial para el seguimiento profesional.")
add(354, "Generar informes, exportar datos e importar pruebas de la plantilla PDF con revisión de los valores detectados.")
add(360, "Valores")
add(362, "El trabajo del equipo se guiará por los valores de Scrum descritos en el modelo:")
for value in (
    "Compromiso con los objetivos acordados para cada entrega.",
    "Coraje para comunicar impedimentos y abordar problemas del producto.",
    "Respeto entre quienes aportan criterios clínicos y técnicos.",
    "Foco en el objetivo de cada Sprint.",
    "Apertura para revisar avances, decisiones y resultados.",
):
    add(363, value)

add(369, "1.2 Roles")
add(371, "Product Owner")
add(373, "Representa las necesidades del profesional y del paciente, ordena el Product Backlog y valida los criterios de aceptación, en especial los relacionados con la interpretación de resultados. La persona que asumirá formalmente este rol queda por confirmar.")
add(374, "Scrum Master")
add(376, "Facilita los eventos Scrum, ayuda a eliminar impedimentos y promueve que el equipo mantenga visible el avance. La persona responsable queda por designar.")
add(377, "Equipo de desarrollo")
add(379, "Diseña, implementa y verifica la aplicación web, las reglas de cálculo, la base de datos y los informes. Sus integrantes y capacidad de trabajo deberán registrarse en la planificación.")

add(388, "1.3 Product Backlog")
add(390, "El Product Backlog reúne las capacidades necesarias para el Test de Líneas. Este resumen toma como base las funciones identificadas en el repositorio; cada elemento deberá detallarse después con historia de usuario, criterio de aceptación y estimación.")
for item in (
    "Acceso profesional y protección de datos por propietario.",
    "Registro, búsqueda y actualización de pacientes.",
    "Creación de enlaces de evaluación asociados a pacientes.",
    "Consentimiento, datos personales y aplicación de la prueba de nueve estímulos.",
    "Cálculo, almacenamiento y consulta de resultados y tiempos.",
    "Historial longitudinal, informes y exportación a Excel.",
    "Importación de pruebas en PDF con revisión manual de los nueve valores.",
):
    add(354, item)

add(397, "1.4 Product Backlog priorizado")
add(399, "La siguiente ordenación es preliminar y debe validarla el Product Owner antes de asignar historias a un Sprint.")
add(400, "Prioridad")
add(402, "Se priorizan primero la confidencialidad, la integridad de la evaluación y la disponibilidad de los resultados; luego, las funciones que facilitan el seguimiento y reducen trabajo manual.")
add(403, "Técnica MoSCoW:")
add(405, "La técnica MoSCoW agrupa los elementos según su necesidad para la primera entrega. Esta clasificación podrá cambiar durante el refinamiento del Product Backlog.")
for item in (
    "M — Necesario: acceso profesional, pacientes, enlaces válidos, consentimiento, prueba completa, cálculo y almacenamiento seguro.",
    "S — Importante: seguimiento longitudinal, informe profesional, exportación e importación revisada desde PDF.",
    "C — Conveniente: mejoras de consulta y presentación basadas en la experiencia de uso.",
    "W — Fuera de este alcance inicial: facturación, teleconsulta e integración con historias clínicas externas.",
):
    add(354, item)

output.save(output_path)
print(output_path)

