from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

root = Path(r"C:\Users\akhse\Desktop\REPO\TestLineas")
reference = Path(r"C:\Users\akhse\Downloads\Metodologia SCRUM - Modelo 1 (1).docx")
output = root / "docs" / "Metodologia_Scrum_Test_de_Lineas_Base.docx"
doc = Document(reference)
source = list(doc.paragraphs)
heading1, heading2, heading3, body_style = (source[i].style for i in (345, 32, 347, 349))

def replace_cover(index, value):
    p = source[index]
    p.runs[0].text = value
    for run in p.runs[1:]:
        run.text = ""

replace_cover(10, "ANÁLISIS, DISEÑO Y DESARROLLO DEL SISTEMA WEB")
replace_cover(12, "Sistema web para la aplicación y seguimiento del Test de Líneas")
replace_cover(16, "Axel Kevin Huaman Solis")
replace_cover(30, "LIMA – PERÚ 2026")

body = doc.element.body
cover_end = source[31]._element
children = list(body)
for child in children[children.index(cover_end) + 1:]:
    if child.tag.endswith("}sectPr"):
        continue
    body.remove(child)

def next_page():
    previous = doc.sections[-1]
    geometry = (previous.page_width, previous.page_height, previous.top_margin,
                previous.bottom_margin, previous.left_margin, previous.right_margin)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    (section.page_width, section.page_height, section.top_margin,
     section.bottom_margin, section.left_margin, section.right_margin) = geometry

def add_index(title):
    p = doc.add_paragraph(style=heading2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    doc.add_paragraph(style=body_style)
    next_page()

add_index("ÍNDICE DE CONTENIDO")
add_index("ÍNDICE DE TABLAS")
add_index("ÍNDICE DE FIGURAS")

def divider(title, continue_after=True):
    p = doc.add_paragraph(style=heading1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(250)
    p.paragraph_format.space_after = Pt(0)
    p.add_run(title)
    if continue_after:
        next_page()

divider("INICIO")
p = doc.add_paragraph(style=heading2)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("1 Alcance del producto")

sections = [
 ("1.1 Propósito del sistema",
  "Test de Líneas es una plataforma web para aplicar una evaluación visual de nueve estímulos. El profesional invita al paciente mediante un enlace; el paciente acepta participar, registra sus datos y ordena las imágenes según su preferencia. El sistema calcula y conserva los resultados para su revisión profesional."),
 ("1.2 Usuarios y partes interesadas",
  "El paciente realiza la prueba desde el enlace recibido. El profesional inicia sesión, administra fichas de pacientes, genera invitaciones y consulta resultados, informes e historial. Las personas que asumirán Product Owner, Scrum Master y equipo de desarrollo deberán identificarse antes de completar los apartados de roles y planificación."),
 ("1.3 Alcance funcional actual",
  "El sistema registra pacientes, permite crear enlaces de evaluación y solicita consentimiento antes de iniciar. Admite ordenar los nueve estímulos, conservar temporalmente el avance y registrar tiempos de respuesta. Calcula interpretaciones, guarda evaluaciones, muestra la evolución del paciente y permite imprimir informes o exportar datos a Excel. También importa pruebas en PDF de la plantilla prevista: detecta los nueve valores y exige revisarlos antes de guardar."),
 ("1.4 Límites del producto",
  "El resultado automático sirve de apoyo y no constituye por sí mismo un diagnóstico. El paciente no ve la interpretación en su dispositivo. La importación de PDF se limita a la plantilla esperada de una página y a archivos de hasta 10 MB. El alcance identificado no incluye facturación, teleconsulta ni integración con historias clínicas externas.")
]
for title, summary in sections:
    p = doc.add_paragraph(style=heading3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(title)
    p = doc.add_paragraph(style=body_style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(9)
    p.add_run(summary)

next_page()
divider("PLANIFICACIÓN Y ESTIMACIÓN")
divider("IMPLEMENTACIÓN", continue_after=False)
doc.paragraphs[32].paragraph_format.page_break_before = True
doc.core_properties.author = "Axel Kevin Huaman Solis"
doc.core_properties.last_modified_by = "Axel Kevin Huaman Solis"
doc.core_properties.title = "Metodología Scrum para Test de Líneas"
output.parent.mkdir(parents=True, exist_ok=True)
doc.save(output)
print(output)



