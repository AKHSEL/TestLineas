from copy import deepcopy
from pathlib import Path
from docx import Document

reference = Document(r"C:\Users\akhse\Downloads\Metodologia SCRUM - Modelo 1 (1).docx")
output_path = Path(r"C:\Users\akhse\Desktop\REPO\TestLineas\docs\Metodologia_Scrum_Test_de_Lineas_Base.docx")
output = Document(output_path)

def set_element_text(element, value):
    texts = element.xpath(".//w:t")
    if not texts:
        raise ValueError("No text node in paragraph or cell")
    texts[0].text = value
    for text in texts[1:]:
        text.text = ""

descriptions = {
    "Representa las necesidades del profesional": "El Dr. Jose Luis Pereyra asume el rol de Product Owner. Define las necesidades del producto, ordena el Product Backlog y valida los criterios de aceptación, especialmente los relacionados con la interpretación profesional de los resultados.",
    "Facilita los eventos Scrum": "Anthoni Rivera asume el rol de Scrum Master. Facilita los eventos Scrum, ayuda a eliminar impedimentos y promueve que el equipo mantenga visible el avance del trabajo.",
    "Diseña, implementa y verifica": "Axel Kevin Huaman Solis integra el Development Team. Diseña, implementa y verifica la aplicación web, las reglas de cálculo, la base de datos y los informes.",
}
for paragraph in output.paragraphs:
    for prefix, replacement in descriptions.items():
        if paragraph.text.startswith(prefix):
            set_element_text(paragraph._element, replacement)

anchor = next(p._element for p in output.paragraphs if p.text == "1.3 Product Backlog")
body = output.element.body

def insert_paragraph(index):
    element = deepcopy(reference.paragraphs[index]._element)
    body.insert(body.index(anchor), element)

def insert_table(index):
    element = deepcopy(reference.tables[index]._element)
    rows = element.xpath("./w:tr")
    from docx.oxml.ns import qn
    second = rows[2].xpath("./w:tc")[1]
    for merge in second.xpath("./w:tcPr/w:vMerge"):
        second.xpath("./w:tcPr")[0].remove(merge)
    last = rows[3].xpath("./w:tc")[1]
    rows[3].replace(last, deepcopy(rows[1].xpath("./w:tc")[1]))
    names = ("Dr. Jose Luis Pereyra", "Anthoni Rivera", "Axel Kevin Huaman Solis")
    for row, name in zip(rows[1:], names):
        cell = row.xpath("./w:tc")[1]
        set_element_text(cell, name)
    body.insert(body.index(anchor), element)

insert_paragraph(380)
insert_paragraph(381)
insert_table(0)
insert_paragraph(382)
insert_paragraph(383)
insert_paragraph(384)
insert_paragraph(385)
insert_table(1)
insert_paragraph(386)
insert_paragraph(387)

output.save(output_path)
print(output_path)


